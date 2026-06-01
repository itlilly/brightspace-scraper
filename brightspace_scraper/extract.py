"""Turn downloaded files into plain text.

Text-bearing formats (PDF/Word/PPT/HTML/txt) are extracted to text here, in Stage 1,
with no model and no OCR — harvesting stays deterministic and dependency-light. Images
and image-only/scanned PDFs are *not* read here; they're flagged `needs_vision` and the
raw file is preserved (the harvester records its path as `content_ref`). The interpreter
side reads those bytes — OCR first, a vision model later — so that work lives wherever
Stage 2 runs (e.g. the GPU desktop) and a thin client never needs a Tesseract binary.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from .util import html_to_text


@dataclass
class Extracted:
    text: str | None
    needs_vision: bool = False


_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp", ".svg"}
_HTML_EXTS = {".html", ".htm"}
_TEXT_EXTS = {".txt", ".md", ".csv"}
# Anything we can't read as text and shouldn't pretend to.
_OPAQUE_EXTS = {".zip", ".mp4", ".mov", ".mp3", ".wav", ".m4a", ".avi"}


def _pdf_text(path: Path) -> str | None:
    text = ""
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            text = "\n".join(p.extract_text() or "" for p in pdf.pages).strip()
    except Exception:
        text = ""
    if text:
        return text
    # Fallback engine for PDFs pdfplumber chokes on.
    try:
        import fitz  # PyMuPDF

        with fitz.open(path) as doc:
            text = "\n".join(page.get_text() for page in doc).strip()
    except Exception:
        text = ""
    return text or None


def _docx_text(path: Path) -> str | None:
    try:
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(x for x in parts if x).strip() or None
    except Exception:
        return None


def _pptx_text(path: Path) -> str | None:
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        parts: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(x for x in parts if x).strip() or None
    except Exception:
        return None


def extract_text(path: Path) -> Extracted:
    """Extract text from a downloaded file based on its extension.

    Images and text-less (scanned) PDFs return `needs_vision=True` with no text — the
    interpreter side OCRs/vision-reads them from the preserved raw file.
    """
    ext = path.suffix.lower()

    if ext in _IMAGE_EXTS:
        return Extracted(text=None, needs_vision=True)
    if ext in _OPAQUE_EXTS:
        return Extracted(text=None, needs_vision=False)

    if ext == ".pdf":
        text = _pdf_text(path)
        # A PDF with no extractable text is almost certainly scanned -> needs_vision.
        return Extracted(text=text, needs_vision=text is None)
    if ext == ".docx":
        return Extracted(text=_docx_text(path))
    if ext == ".pptx":
        return Extracted(text=_pptx_text(path))
    if ext in _HTML_EXTS:
        raw = path.read_text(encoding="utf-8", errors="replace")
        return Extracted(text=html_to_text(raw))
    if ext in _TEXT_EXTS or not ext:
        raw = path.read_text(encoding="utf-8", errors="replace").strip()
        return Extracted(text=raw or None)

    return Extracted(text=None)
